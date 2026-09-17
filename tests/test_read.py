import openpyxl
from openpyxl.styles import PatternFill
import pytest

from ftmap.io.detect import sniff_delimiter, sniff_encoding
from ftmap.io.tabular import _cell, read_source, sha256_file


def test_cp1251_csv_is_read_without_mojibake(tmp_path):
    p = tmp_path / "a.csv"
    p.write_bytes("ПІБ,Посада\nКоваленко,депутат\n".encode("cp1251"))
    assert sniff_encoding(str(p)) == "cp1251"
    grids = read_source(str(p))
    assert grids[0].rows[0] == ["ПІБ", "Посада"]


def test_semicolon_csv_with_commas_inside_values(tmp_path):
    p = tmp_path / "b.csv"
    p.write_text(
        "Посада;ПІБ;Час\nГолова ради;Стецюк Володимир;1-й четвер, 10:00\n",
        encoding="utf-8",
    )
    assert sniff_delimiter(p.read_text(encoding="utf-8")) == ";"
    grids = read_source(str(p))
    assert grids[0].rows[1][2] == "1-й четвер, 10:00"


def test_excel_sheets_become_separate_grids_and_merges_are_reported(tmp_path):
    p = tmp_path / "c.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "перший"
    ws["A1"] = "СТРУКТУРА"
    ws.merge_cells("A1:C1")
    ws.append([])
    ws.append(["№ з/п", "Назва", "Штат"])
    wb.create_sheet("другий").append(["x"])
    wb.save(p)
    grids = read_source(str(p))
    assert [g.sheet for g in grids] == ["перший", "другий"]
    assert grids[0].merges == [(0, 0, 0, 2)]
    assert grids[0].rows[2] == ["№ з/п", "Назва", "Штат"]


def test_an_integral_float_cell_is_read_as_the_integer_it_is():
    """xlrd hands every numeric cell over as a float, so a year typed as 2006
    arrives as `2006.0` — and a date canonicalizer refuses it, an identifier
    canonicalizer keeps the `.0`, and the model, shown `2006.0`, declines
    `buildDate` for a column of years. Measured on the aircraft register: 872
    of 872 build years shaped `dddd.d`, acceptance 0.00. A float that IS an
    integer is written as one; a real fraction is left alone."""
    assert _cell(2006.0) == "2006"
    assert _cell(-3.0) == "-3"
    assert _cell(3.5) == "3.5"
    assert _cell(2006) == "2006"
    assert _cell(True) == "True"
    assert _cell(float("nan")) == "nan"


def test_sha256_is_stable(tmp_path):
    p = tmp_path / "d.csv"
    p.write_text("a,b\n1,2\n", encoding="utf-8")
    assert sha256_file(str(p)) == sha256_file(str(p))
    assert len(sha256_file(str(p))) == 64


def test_xls_named_file_that_is_really_xlsx_falls_back(tmp_path):
    # Common in the real corpus: a workbook saved as XLSX but named .xls.
    # xlrd 2.x cannot open it (XLSX support was dropped), so read_source
    # must catch that failure and retry with the openpyxl-based reader.
    p = tmp_path / "e.xls"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "аркуш"
    ws.append(["Прізвище", "Посада"])
    ws.append(["Коваленко", "депутат"])
    wb.save(p)
    grids = read_source(str(p))
    assert grids[0].sheet == "аркуш"
    assert grids[0].rows[1] == ["Коваленко", "депутат"]


def test_blank_row_is_not_dropped(tmp_path):
    # A reader that silently skipped a blank row would shift every later row
    # index, and provenance in the final output would then point at the
    # wrong cell. The blank line must survive as a row of its own.
    p = tmp_path / "f.csv"
    p.write_text("a,b\n1,2\n\n3,4\n", encoding="utf-8")
    grids = read_source(str(p))
    assert len(grids[0].rows) == 4
    assert grids[0].rows[2] == []
    assert grids[0].rows[3] == ["3", "4"]


def test_formatting_only_columns_past_the_data_are_not_read_as_columns(tmp_path):
    """A style applied to an unused column extends openpyxl's `max_column`, and
    `iter_rows` pads every row out to it. Measured on `data/files`: one sheet
    reports 2 575 columns where 34 hold anything, another reports Excel's full
    16 384 where 30 do — 39 million cells built, held and profiled for thirty
    real ones. Trailing padding is not data the file has, and dropping it moves
    no surviving column, because every one of them is to its left."""
    p = tmp_path / "padded.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ПІБ", "Посада"])
    ws.append(["Коваленко Іван Петрович", "депутат"])
    ws.cell(row=1, column=300).fill = PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    wb.save(p)
    assert openpyxl.load_workbook(p).worksheets[0].max_column == 300
    grid = read_source(str(p))[0]
    assert [len(r) for r in grid.rows] == [2, 2]


def test_an_empty_column_between_two_populated_ones_is_kept(tmp_path):
    """Only TRAILING padding goes. An empty column inside the data is a column
    the file really has, and dropping it would shift every column to its right
    — which is the provenance the row-keeping contract above exists to
    protect, one axis over."""
    p = tmp_path / "gap.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ПІБ", None, "Посада"])
    ws.append(["Коваленко Іван Петрович", None, "депутат"])
    wb.save(p)
    grid = read_source(str(p))[0]
    assert [len(r) for r in grid.rows] == [3, 3]
    assert grid.rows[0] == ["ПІБ", None, "Посада"]


def test_a_trailing_column_that_holds_only_a_header_is_kept(tmp_path):
    """A header with no values under it is a column the source declares and
    `profile` reports as unfilled. Only a column with nothing anywhere in it,
    header included, is padding."""
    p = tmp_path / "header-only.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ПІБ", "Посада", "Примітка"])
    ws.append(["Коваленко Іван Петрович", "депутат", None])
    wb.save(p)
    grid = read_source(str(p))[0]
    assert [len(r) for r in grid.rows] == [3, 3]
    assert grid.rows[0][2] == "Примітка"


def test_a_merge_reaching_past_the_trimmed_width_is_clamped(tmp_path):
    """A title merged across the formatted extent must not survive as a merge
    over columns that no longer exist: `layout` fills merged cells forward, and
    a merge past the last column would fill nothing at an index no row has."""
    p = tmp_path / "wide-merge.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "СТРУКТУРА"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=200)
    ws.append(["№ з/п", "Назва"])
    wb.save(p)
    grid = read_source(str(p))[0]
    width = len(grid.rows[0])
    assert width == 2
    assert all(c1 < width for _, _, _, c1 in grid.merges)


def test_a_utf8_file_cut_mid_character_by_the_sampler_is_still_utf8(tmp_path):
    """The sampler reads a fixed 200 000 bytes, and a Cyrillic file whose cut
    lands inside a two-byte character raised `UnicodeDecodeError` for the
    TRUNCATION rather than for the file. cp1251 decodes any byte sequence, so
    it won the fallback and every value came out as mojibake — silently,
    because nothing after this ever sees the bytes again.

    Measured on `corpus-external/person-graph/ua_land_valuers_register.csv`
    on 2026-08-29: a clean UTF-8 file, read as cp1251, 2 892 rows of names
    like `Антипенко Ірина Вікторівна` emitted as
    `РђРЅС‚РёРїРµРЅРєРѕ Р†СЂРёРЅР° Р’С–РєС‚РѕСЂС–РІРЅР°`.
    """
    p = tmp_path / "big.csv"
    row = "Коваленко Іван Петрович,Сумська область\n"
    # THE CUT HAS TO LAND INSIDE A CHARACTER or the test proves nothing, and
    # where it lands depends on the header's length — so the header is padded
    # until it does.
    for pad in range(4):
        text = "ім'я,область" + "x" * pad + "\n" + row * 12_000
        raw = text.encode("utf-8")[:200_000]
        try:
            raw.decode("utf-8")
        except UnicodeDecodeError:
            break
    else:
        raise AssertionError("no padding put a character across the 200 000th byte")
    p.write_text(text, encoding="utf-8")
    # Either UTF-8 spelling is right — `utf-8-sig` differs only in stripping a
    # BOM this file does not have. What must not happen is cp1251, which
    # decodes the same bytes into mojibake without failing.
    encoding = sniff_encoding(str(p))
    assert encoding in ("utf-8", "utf-8-sig")
    assert p.read_text(encoding=encoding).startswith("ім'я")


def test_a_small_file_cut_mid_character_is_readable_with_its_answer(tmp_path):
    """The whole file fits in the sniffer's sample, so a trailing half
    character is the FILE's truncation, not the sampler's. `final=False`
    treated it as the sampler's: the file sniffed as UTF-8 and the strict
    full read then raised where the old cp1251 fallback had processed it."""
    p = tmp_path / "cut.csv"
    p.write_bytes("ПІБ\nКоваленко Іван".encode("utf-8")[:-1])
    enc = sniff_encoding(str(p))
    with open(str(p), encoding=enc, newline="") as fh:
        fh.read()  # must not raise: the sniff answer must be one that decodes


def test_a_header_delimited_unlike_its_body_is_split_on_its_own_delimiter(tmp_path):
    """The Ministry of Justice debtors register: `A;B;C` over comma-delimited
    quoted rows. The sniffer picks the body's comma and the header came
    through as one cell, which the layout then read as a title."""
    from ftmap.io.tabular import read_source
    p = tmp_path / "erb.csv"
    p.write_text('DEBTOR_NAME;DEBTOR_CODE;VD_CAT\r\n'
                 '"ТОВ ""Денкар""","32841550","стягнення коштів"\r\n'
                 '"Коваленко Іван","","аліменти, борг"\r\n', encoding="cp1251")
    g = read_source(str(p))[0]
    assert g.rows[0] == ["DEBTOR_NAME", "DEBTOR_CODE", "VD_CAT"]
    assert g.rows[2] == ["Коваленко Іван", None, "аліменти, борг"]
    assert g.repairs and "delimited by ';'" in g.repairs[0]


def test_a_one_cell_title_row_over_a_wide_body_is_not_split(tmp_path):
    from ftmap.io.tabular import read_source
    p = tmp_path / "t.csv"
    p.write_text("Реєстр суден; станом на 1 липня\nномер,назва,власник\n1,Аврора,Коваленко\n",
                 encoding="utf-8")
    g = read_source(str(p))[0]
    assert g.rows[0][0].startswith("Реєстр") and len([c for c in g.rows[0] if c]) == 1
    assert g.repairs == []
