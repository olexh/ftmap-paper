# tests/test_inventory.py
"""`source_id` is `sha256[:12]/<sheet>` precisely so identical content is one
source (see `inventory.py`'s module docstring). These tests are the guard on
that claim: C2 found that the code built one `SourceRef` per PATH and never
deduplicated, so the same file under two paths — a copy, a symlink target
walked twice — was counted as two sources."""

import openpyxl
import pytest

from ftmap.inventory import SourceRef, _path_prefix, inventory


def test_two_copies_of_one_file_share_a_source_id_and_only_one_is_kept(tmp_path):
    """Measured on `article-1/01-data/raw_835`: 11 duplicate-content pairs
    among 355 files. Every ref is still returned — this is a reported fact,
    not a silent merge — but only one of them is unmarked."""
    d = tmp_path / "corpus"
    d.mkdir()
    content = "ПІБ,Дата\nКоваленко Іван,17.09.1980\n"
    (d / "a.csv").write_text(content, encoding="utf-8")
    (d / "b_copy.csv").write_text(content, encoding="utf-8")

    refs = inventory(str(d))
    assert len(refs) == 2
    assert len({r.source_id for r in refs}) == 1

    kept = [r for r in refs if r.duplicate_of is None]
    dup = [r for r in refs if r.duplicate_of is not None]
    assert len(kept) == 1 and len(dup) == 1
    # "a.csv" sorts before "b_copy.csv", so the first path in sorted order is kept.
    assert kept[0].path.endswith("a.csv")
    assert dup[0].path.endswith("b_copy.csv")
    assert dup[0].duplicate_of == kept[0].path


def test_distinct_content_is_never_marked_duplicate(tmp_path):
    d = tmp_path / "corpus"
    d.mkdir()
    (d / "a.csv").write_text("ПІБ\nКоваленко Іван\n", encoding="utf-8")
    (d / "b.csv").write_text("ПІБ\nШевченко Ольга\n", encoding="utf-8")

    refs = inventory(str(d))
    assert all(r.duplicate_of is None for r in refs)


def test_a_workbook_with_two_sheets_is_two_sources_not_a_duplicate(tmp_path):
    """Two sheets of one file describe different things and must never be
    folded together, even though they share one `sha256`."""
    import openpyxl

    d = tmp_path / "corpus"
    d.mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "перший"
    wb.active.append(["ПІБ"])
    wb.create_sheet("другий").append(["Назва"])
    wb.save(d / "book.xlsx")

    refs = inventory(str(d))
    assert len(refs) == 2
    assert all(r.duplicate_of is None for r in refs)
    assert len({r.source_id for r in refs}) == 2


def _workbook(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(title=name)
        ws.append(["ПІБ"])
        ws.append(["Коваленко Іван"])
    wb.save(path)


def test_sheet_names_that_sanitize_alike_still_get_different_paths(tmp_path):
    """`safe_id` was a lossy rewrite of `source_id`, and lossy means colliding.

    It replaced `/` with `__` and a space with `_` and nothing else, so a
    workbook with sheets `a b` and `a_b` produced ONE directory for two
    sources: the second run overwrote the first's profile, plan, entities and
    summary, and the corpus roll-up counted one of them twice. Excel allows
    both names in one file, which is all it takes.
    """
    path = str(tmp_path / "w.xlsx")
    _workbook(path, ["a b", "a_b"])
    refs = [r for r in inventory(path) if r.duplicate_of is None]
    assert len({r.source_id for r in refs}) == 2
    assert len({r.safe_id for r in refs}) == 2, [r.safe_id for r in refs]


def test_a_hostile_sheet_name_becomes_a_readable_path(tmp_path):
    """Sanitized for the filesystem AND for followthemoney's mapping loader.

    `#`, `?` and `$` each make that loader read a different file than the
    mapping names — see `plan/compile.UNSAFE_URL_CHARS` for the measurement —
    and a sheet name is attacker-supplied in exactly the way a workbook
    downloaded from an open-data portal is attacker-supplied.

    Excel's own rules forbid `? / \\ : * [ ]` in a sheet title and openpyxl
    enforces them, so a name carrying one cannot be WRITTEN here — which is
    exactly why the sanitizer is tested against the raw string as well as
    through a workbook. A crafted file is not obliged to obey Excel.
    """
    hostile = "Звіт #1 $HOME? a/b\tc:d"
    prefix = _path_prefix(f"abc123def456/{hostile}")
    assert not (set(prefix) & set("#?$/\\:\t\r\n")), prefix
    # Readable, not opaque: the point of a prefix is that an operator can find
    # the directory for the sheet they are looking at.
    assert "Звіт" in prefix

    # And through a real workbook, with the subset Excel actually permits.
    path = str(tmp_path / "w.xlsx")
    _workbook(path, ["Звіт #1 $HOME"])
    ref = inventory(path)[0]
    assert not (set(ref.safe_id) & set("#?$/\\:\t\r\n"))
    assert "Звіт" in ref.safe_id


def test_a_long_sheet_name_is_cut_without_losing_identity(tmp_path):
    """A path component is 255 bytes, and Cyrillic is two bytes a character.

    Truncation is safe only because identity lives in the suffix, so this
    asserts both halves: the name fits, and two names that differ only past
    the cut still get two directories.
    """
    long_a = "Ф" * 200 + "A"
    long_b = "Ф" * 200 + "B"
    ids = {SourceRef(path="/x.xlsx", sheet_index=i, sheet=name, sha256="0" * 64,
                     size=1, ext=".xlsx", source_id=f"abc123def456/{name}").safe_id
           for i, name in enumerate((long_a, long_b))}
    assert len(ids) == 2
    assert all(len(i.encode("utf-8")) < 255 for i in ids)


def test_the_inventory_refuses_to_run_if_two_sources_share_a_path(
        tmp_path, monkeypatch):
    """The assertion that makes the encoding's injectivity checkable.

    A digest suffix makes a collision vanishingly unlikely rather than
    impossible, and "unlikely" is not a property a run can rely on when the
    consequence is one directory holding two sources. Forced here by making
    the suffix constant, which is the only way to reach the branch.
    """
    import ftmap.inventory as inv

    path = str(tmp_path / "w.xlsx")
    _workbook(path, ["звіт a", "звіт_a"])
    monkeypatch.setattr(inv, "_path_suffix", lambda source_id: "deadbeef")
    with pytest.raises(inv.SourcePathCollision) as excinfo:
        inventory(path)
    assert "звіт a" in str(excinfo.value)
