# tests/test_pipeline.py
import json
import os

import followthemoney
import openpyxl
import pytest

from ftmap.config import Config
from ftmap.inventory import inventory
from ftmap.pipeline import (PublishRefused, _apply_overrides, run_corpus,
                            run_source)
from ftmap.plan.compile import NO_ENTITY_REASON
from ftmap.plan.prompt import BINDING_MARKER
from ftmap.plan.validate import ValidatedPlan
from ftmap.vocab.catalogue import Catalogue
from ftmap.workspace import open_workspace

CFG = Config.load(None)
CAT = Catalogue.load()


def _read(path: str) -> str:
    with open(path, encoding="utf-8") as fh:
        return fh.read()


class ScriptedClient:
    """Structure then bindings, keyed on whether the prompt declares entities."""

    def __init__(self):
        self.calls = 0

    def complete(self, system, user, schema, max_tokens=None):
        self.calls += 1
        if BINDING_MARKER not in user:
            return {"subject": "Person",
                    "entities": [{"key": "person", "schema": "Person", "keys": ["c0"]}],
                    "edges": []}
        out = []
        for line in user.splitlines():
            if line.startswith("c") and ":" in line:
                cid = line.split(":", 1)[0].strip()
                prop = "Person:name" if cid == "c0" else "Person:birthDate"
                out.append({"column": cid, "binding": f"person|{prop}", "why": "x"})
        return {"bindings": out}


# THE LAYOUT MOVED, AND THESE ARE THE THREE QUESTIONS THAT MOVED WITH IT.
# A source's artefacts are no longer at `<out>/<safe_id>/`; they are in the
# committed generation the manifest names, and the corpus manifest is inside
# the committed aggregate. Every test below asks through `open_workspace` for
# the same reason production code does: a path built by hand is a path that
# can name a staging directory, a superseded generation, or a stranger.
def _sdir(out, safe_id):
    return open_workspace(out).source_dir(safe_id)


def _mapping(out, safe_id):
    return open_workspace(out).mapping_path(safe_id)


def _report(out):
    return open_workspace(out).report()


def _corpus(tmp_path):
    d = tmp_path / "corpus"
    d.mkdir()
    (d / "a.csv").write_text("ПІБ,Дата народження\nКоваленко Іван,17.09.1980\n",
                             encoding="utf-8")
    wb = openpyxl.Workbook()
    wb.active.title = "перший"
    wb.active.append(["ПІБ", "Дата народження"])
    wb.active.append(["Шевченко Ольга", "01.02.1990"])
    wb.save(d / "b.xlsx")
    return str(d)


def test_inventory_expands_workbooks_to_one_source_per_sheet(tmp_path):
    refs = inventory(_corpus(tmp_path))
    assert len(refs) == 2
    ids = {r.source_id for r in refs}
    assert all("/" in i for i in ids)
    assert len({r.sha256 for r in refs}) == 2


def test_run_source_writes_every_artefact(tmp_path):
    refs = inventory(_corpus(tmp_path))
    out = str(tmp_path / "work")
    summary = run_source(refs[0], CFG, CAT, ScriptedClient(), "run1", out)
    d = _sdir(out, summary["safe_id"])
    for name in ("profile.json", "plan.json", "plan.validated.json",
                 "normalized.csv", "rejects.jsonl", "decisions.jsonl",
                 "entities.ftm.json", "statements.csv", "summary.json"):
        assert os.path.exists(os.path.join(d, name)), name
    assert os.path.exists(_mapping(out, summary["safe_id"]))


def test_the_run_report_says_which_prompts_produced_it(tmp_path):
    """Three individually-correct edits to the structure prompt produced three
    different valid decompositions of one real file. A run's numbers are only
    interpretable beside the text that produced them."""
    out = str(tmp_path / "work")
    report = run_corpus(_corpus(tmp_path), CFG, CAT, ScriptedClient(), out)
    assert set(report["prompts"]) == {"structure", "edge", "binding"}
    assert all(len(v) == 12 for v in report["prompts"].values())


def test_run_corpus_merges_and_reports_totals(tmp_path):
    out = str(tmp_path / "work")
    report = run_corpus(_corpus(tmp_path), CFG, CAT, ScriptedClient(), out)
    assert report["sources"] == 2
    assert report["duplicates"] == 0
    merged = os.path.join(open_workspace(out).aggregate_dir(), "entities.ftm.json")
    with open(merged, encoding="utf-8") as fh:
        lines = [json.loads(l) for l in fh]
    assert len(lines) == 2
    assert report["entities"]["Person"] == 2


def test_run_id_does_not_collide_across_two_runs_of_the_same_corpus(tmp_path):
    """I7: run_id was built from the corpus basename and the source count
    alone, so two runs over the same corpus produced the SAME run_id — the
    field every statement carries as its own provenance stopped being an
    identifier. A rerun that replaces a batch must be distinguishable from
    the batch it replaced."""
    root = _corpus(tmp_path)
    first = run_corpus(root, CFG, CAT, ScriptedClient(), str(tmp_path / "work1"))
    second = run_corpus(root, CFG, CAT, ScriptedClient(), str(tmp_path / "work2"))
    assert first["run_id"] != second["run_id"]


def test_run_json_records_versions_and_the_start_time(tmp_path):
    """I7: a published measurement should be reproducible from its own
    manifest. The spec fixes ftmap's own version and FollowTheMoney's."""
    out = str(tmp_path / "work")
    report = run_corpus(_corpus(tmp_path), CFG, CAT, ScriptedClient(), out)
    assert report["ftmap_version"] == "0.1.0"
    assert report["followthemoney_version"] == followthemoney.__version__
    assert report["started_at"]  # non-empty ISO 8601 timestamp
    # A DETERMINISTIC RUN SAYS SO, RATHER THAN LEAVING A NULL. `null` meant
    # two things a reader cannot tell apart: no model was ever involved, and
    # a model-backed run whose every call was a cache hit. The first is this
    # one, and it now states itself.
    assert report["model_id"] == "not_applicable"
    assert report["model_revision"] == "not_applicable"


def test_run_corpus_processes_a_duplicate_file_once_and_reports_it(tmp_path):
    """C2: the same file under two paths must be counted and processed as one
    source. Before the fix, `columns.total` and `statements.csv` — the
    coverage numbers this project publishes — double-counted every duplicate,
    while `entities.ftm.json` happened to survive via id-based dedup."""
    d = tmp_path / "corpus"
    d.mkdir()
    content = "ПІБ,Дата народження\nКоваленко Іван,17.09.1980\n"
    (d / "a.csv").write_text(content, encoding="utf-8")
    (d / "a_copy.csv").write_text(content, encoding="utf-8")
    out = str(tmp_path / "work")

    report = run_corpus(str(d), CFG, CAT, ScriptedClient(), out)
    assert report["sources"] == 1
    assert report["duplicates"] == 1
    assert report["duplicate_sources"][0]["path"].endswith("a_copy.csv")
    assert report["duplicate_sources"][0]["duplicate_of"].endswith("a.csv")

    with open(os.path.join(open_workspace(out).aggregate_dir(), "statements.csv"), encoding="utf-8") as fh:
        stmt_lines = fh.readlines()
    assert len(stmt_lines) == 3  # header + name + birthDate, from ONE source


def test_every_value_bearing_artefact_is_owner_only(tmp_path):
    """The model's free-text `why` was measured to quote real cell values in 13
    of 148 bindings on a real run, and plan.json, plan.validated.json and
    decisions.jsonl were left world-readable because nobody had checked."""
    refs = inventory(_corpus(tmp_path))
    out = str(tmp_path / "work")
    summary = run_source(refs[0], CFG, CAT, ScriptedClient(), "run1", out)
    d = _sdir(out, summary["safe_id"])
    for name in ("profile.json", "plan.json", "plan.validated.json",
                 "decisions.jsonl", "normalized.csv", "rejects.jsonl",
                 "entities.ftm.json", "statements.csv"):
        mode = oct(os.stat(os.path.join(d, name)).st_mode)[-3:]
        assert mode == "600", f"{name} is {mode}"
    assert oct(os.stat(os.path.join(d, "summary.json")).st_mode)[-3:] == "644"


def test_a_root_that_cannot_be_walked_is_not_an_empty_corpus(tmp_path):
    """os.walk swallows the OSError by default, so an unreadable root produced
    `sources: 0, failed: 0` — a clean-looking run over nothing."""
    with pytest.raises(FileNotFoundError):
        inventory(str(tmp_path / "does-not-exist"))


def test_a_source_that_raises_is_reported_not_swallowed(tmp_path):
    d = tmp_path / "corpus"
    d.mkdir()
    (d / "broken.xlsx").write_bytes(b"not really a workbook")
    out = str(tmp_path / "work")
    report = run_corpus(str(d), CFG, CAT, ScriptedClient(), out)
    assert report["failed"] == 1
    assert report["failures"][0]["path"].endswith("broken.xlsx")


def test_run_json_failures_carry_no_model_output_only_a_stable_code(tmp_path):
    """I6: run.json is published at 0644. Before the fix, its failures[]
    entries carried the exception's full str() and a formatted traceback —
    and client.py's ModelError embeds up to 200 characters of the model's
    raw answer, whose `why` was separately measured to quote real cell
    values. The full record now lives only in failures.jsonl, mode 0600;
    run.json gets an error_type and a stable code, and nothing else."""
    d = tmp_path / "corpus"
    d.mkdir()
    (d / "broken.xlsx").write_bytes(b"not really a workbook")
    out = str(tmp_path / "work")

    report = run_corpus(str(d), CFG, CAT, ScriptedClient(), out)
    assert oct(os.stat(os.path.join(open_workspace(out).aggregate_dir(), "run.json")).st_mode)[-3:] == "644"
    failure = report["failures"][0]
    assert set(failure) == {"path", "sheet", "error_type", "code"}

    failures_path = os.path.join(open_workspace(out).aggregate_dir(), "failures.jsonl")
    assert oct(os.stat(failures_path).st_mode)[-3:] == "600"
    full = json.loads(_read(failures_path).splitlines()[0])
    assert full["error_type"] == failure["error_type"]
    assert full["code"] == failure["code"]
    assert "error" in full and "traceback" in full

    # Stable: rerunning the same broken source raises from the same site.
    report2 = run_corpus(str(d), CFG, CAT, ScriptedClient(), out)
    assert report2["failures"][0]["code"] == failure["code"]


def test_rerun_is_byte_identical(tmp_path):
    root, out = _corpus(tmp_path), str(tmp_path / "work")
    run_corpus(root, CFG, CAT, ScriptedClient(), out)
    first = _read(os.path.join(open_workspace(out).aggregate_dir(), "entities.ftm.json"))
    run_corpus(root, CFG, CAT, ScriptedClient(), out)
    assert _read(os.path.join(open_workspace(out).aggregate_dir(), "entities.ftm.json")) == first


def test_a_crash_on_rerun_leaves_the_committed_generation_untouched(tmp_path, monkeypatch):
    """I1: `execute()` can raise well after plan.validated.json,
    decisions.jsonl, normalized.csv and the mapping YAML have been written —
    reachable from the review app on a bad override, or any other FtM-level
    rejection of a hand-edited plan. It used to leave the source directory
    carrying two different runs at once: the FRESH (failed)
    plan.validated.json beside the STALE (previous good) entities.ftm.json,
    with no marker that they disagreed.

    THERE IS NOTHING TO ROLL BACK NOW, which is the stronger claim. The
    committed generation was never opened for writing, so the assertions
    below are about a directory the failed attempt could not have reached —
    not about a restore that has to run correctly under a crash."""
    refs = inventory(_corpus(tmp_path))
    out = str(tmp_path / "work")
    good = run_source(refs[0], CFG, CAT, ScriptedClient(), "run1", out)
    d = _sdir(out, good["safe_id"])
    before = {name: _read(os.path.join(d, name))
             for name in ("plan.validated.json", "decisions.jsonl",
                          "entities.ftm.json", "statements.csv", "summary.json")}
    mapping_path = _mapping(out, good["safe_id"])
    mapping_before = _read(mapping_path)

    def boom(*a, **k):
        raise RuntimeError("simulated FtM rejection of the rerun's mapping")

    monkeypatch.setattr("ftmap.pipeline.execute_into", boom)
    with pytest.raises(RuntimeError):
        run_source(refs[0], CFG, CAT, ScriptedClient(), "run2", out)

    for name, content in before.items():
        assert _read(os.path.join(d, name)) == content, name
    assert _read(mapping_path) == mapping_before
    # The manifest still names the same generation, and the failed attempt
    # left no discoverable debris — no rollback copy that a directory scan
    # would read as a second source, and no staging directory either.
    workspace = open_workspace(out)
    assert workspace.source_ids() == [good["safe_id"]]
    assert workspace.source_dir(good["safe_id"]) == d
    assert not any(".rollback-" in n for n in os.listdir(out))
    staging = os.path.join(out, ".staging")
    assert not os.path.isdir(staging) or os.listdir(staging) == []


def _vplan():
    return ValidatedPlan(
        subject="Person",
        entities=[{"key": "person", "schema": "Person", "keys": ["c0"]}],
        edges=[], bindings=[
            {"column": "c0", "prop": "Person:name", "entity": "person",
             "type_name": "name", "why": "names"},
        ])


def test_an_override_naming_a_property_the_entity_does_not_carry_is_rejected():
    """I2: `_apply_overrides` checked only that the qname existed, not that
    the entity's schema carries it. `Membership:role` reaching `person`
    (schema `Person`) used to sail through as "accepted" and crash
    compile/execute instead of being rejected here."""
    v = _apply_overrides(_vplan(), {
        "c1": {"prop": "Membership:role", "entity": "person", "why": "x"},
    }, CAT)
    assert [b["column"] for b in v.bindings] == ["c0"]
    d = [d for d in v.decisions if d.column == "c1"][0]
    assert d.verdict == "rejected" and "does not carry" in d.reason


def test_an_override_naming_an_undeclared_entity_is_rejected_not_silently_dropped():
    """Before the fix, this bypassed every check, was appended to `bindings`
    with entity "ghost", and recorded an "accepted" Decision — the binding
    then vanished at compile time with no trace of why, while the decision
    log said it had succeeded."""
    v = _apply_overrides(_vplan(), {
        "c1": {"prop": "Person:name", "entity": "ghost", "why": "x"},
    }, CAT)
    assert [b["column"] for b in v.bindings] == ["c0"]
    d = [d for d in v.decisions if d.column == "c1"][0]
    assert d.verdict == "rejected" and "not declared" in d.reason


def test_an_override_with_no_declared_entities_and_no_entity_named_does_not_crash():
    """`vplan.entities[0]["key"]` raised IndexError outright when the source
    declared no entities. A rejection Decision, not a stack trace."""
    empty = ValidatedPlan(subject=None, entities=[], edges=[], bindings=[])
    v = _apply_overrides(empty, {
        "c0": {"prop": "Person:name", "why": "x"},
    }, CAT)
    assert v.bindings == []
    d = [d for d in v.decisions if d.column == "c0"][0]
    assert d.verdict == "rejected"


def test_an_override_onto_a_property_the_model_already_bound_takes_it_over():
    """The analyst outranks the model, and one property still reads one
    column: without this the override and the model binding both reached
    `compile_mapping`, which kept whichever it saw last."""
    v = _apply_overrides(_vplan(), {
        "c1": {"prop": "Person:name", "entity": "person", "why": "the real name"},
    }, CAT)
    assert [b["column"] for b in v.bindings] == ["c1"]
    d = [d for d in v.decisions if d.column == "c0"][-1]
    assert d.verdict == "unmapped" and d.decided_by == "analyst"
    assert "the analyst bound Person:name to c1" in d.reason


class TwoColumnsOneProperty:
    """Binds every column of the source to `Person:name` — two columns, one
    property. A name is a SET in followthemoney, so both columns feed it
    (`plan/claims.py`); the МВС registry's four figures on `Vehicle:amount`,
    the shape the accounting was written for, is a `number` and still
    contests — `tests/test_validate.py` and `tests/test_execute.py` pin that
    half on dates."""

    def complete(self, system, user, schema, max_tokens=None):
        if BINDING_MARKER not in user:
            return {"subject": "Person",
                    "entities": [{"key": "person", "schema": "Person", "keys": ["c0"]}],
                    "edges": []}
        return {"bindings": [
            {"column": line.split(":", 1)[0].strip(),
             "binding": "person|Person:name", "why": "x"}
            for line in user.splitlines() if line.startswith("c") and ":" in line
        ]}


def test_the_column_that_shares_a_property_is_bound_not_lost(tmp_path):
    """The whole point of the accounting: every value of both columns is
    offered to the engine, each column reads as bound, and nothing lands in
    `unaccounted`, which is the alarm for a silent loss."""
    refs = inventory(_corpus(tmp_path))
    out = str(tmp_path / "work")
    summary = run_source(refs[0], CFG, CAT, TwoColumnsOneProperty(), "run1", out)
    cov = summary["coverage"]
    assert cov["c0"]["status"] == "bound"
    assert cov["c1"]["status"] == "bound"
    assert cov["c1"]["values"] == cov["c1"]["offered"] > 0
    assert all(c["unaccounted"] == 0 for c in cov.values())
    assert summary["coverage_totals"]["unaccounted"] == 0
    assert summary["coverage_totals"]["accounting_valid"] is True
    shared = [d for d in _decisions(summary) if d["column"] == "c1"
              and d["decided_by"] == "rule" and d["verdict"] == "accepted"]
    assert shared and "c0 also binds Person:name" in shared[-1]["reason"]


def _decisions(summary):
    import json as _json
    path = os.path.join(os.path.dirname(summary["mapping"]), "decisions.jsonl")
    with open(path, encoding="utf-8") as fh:
        return [_json.loads(line) for line in fh if line.strip()]


def test_a_shared_property_is_written_as_one_multi_column_source(tmp_path):
    """One property, one source, two columns: FollowTheMoney's own
    `columns: [...]`, which `ftm map` reads standalone."""
    refs = inventory(_corpus(tmp_path))
    out = str(tmp_path / "work")
    summary = run_source(refs[0], CFG, CAT, TwoColumnsOneProperty(), "run1", out)
    with open(summary["mapping"], encoding="utf-8") as fh:
        text = fh.read()
    assert "columns:" in text
    assert "- c0" in text and "- c1" in text
    assert "column: c1" not in text


class NoEntitySurvives:
    """The МВС vehicle-registry shape, offline: a subject the plan's only
    declared entity contradicts, and a binding call that accepts nothing.

    `resolve_subject_contradiction` declares the subject beside the entity as
    a second reading, the binding call uses neither, `validate` drops both —
    and the plan that reaches `compile_mapping` declares no entity at all.
    """

    calls = 0

    def complete(self, system, user, schema, max_tokens=None):
        if BINDING_MARKER not in user:
            return {"subject": "Vehicle",
                    "entities": [{"key": "org", "schema": "Organization",
                                  "keys": ["c0"]}]}
        return {"bindings": []}


def test_a_source_whose_plan_declares_no_entity_is_declined_not_failed(tmp_path):
    """Measured on the private corpus: 20 of 244 sources compiled
    `{"entities": {}}`, which `QueryMapping.__init__` refuses, and raised out
    of `execute`. A source that raises writes no summary, so those 20 were
    absent from `run.json`'s `sources` count AND from every coverage total the
    run publishes — not reported as lost, just gone."""
    out = str(tmp_path / "work")
    report = run_corpus(_corpus(tmp_path), CFG, CAT, NoEntitySurvives(), out)

    assert report["failed"] == 0 and report["failures"] == []
    assert report["sources"] == 2 and report["declined"] == 2
    assert {s["safe_id"] for s in report["declined_sources"]} == set(
        report["summaries"])

    # THE COLUMNS ARE STILL IN THE DENOMINATOR, which is the point.
    assert report["columns"]["total"] == 4 and report["columns"]["mapped"] == 0
    assert report["statements"] == 0
    corpus = report["coverage"]
    assert corpus["columns"] == 4
    assert corpus["values"] > 0
    assert corpus["declined"] == corpus["values"]
    assert corpus["offered"] == 0 and corpus["unaccounted"] == 0


def test_a_declined_source_says_why_in_summary_and_in_decisions(tmp_path):
    refs = inventory(_corpus(tmp_path))
    out = str(tmp_path / "work")
    summary = run_source(refs[0], CFG, CAT, NoEntitySurvives(), "run1", out)
    d = _sdir(out, summary["safe_id"])

    assert summary["declined"] == NO_ENTITY_REASON
    assert summary["entities"] == {} and summary["statements"] == 0
    # The reason is in decisions.jsonl too, as a verdict about the SOURCE:
    # every other line there is about a column, an entity key or an edge.
    lines = [json.loads(l) for l in _read(os.path.join(d, "decisions.jsonl"))
             .splitlines()]
    declines = [l for l in lines if l["verdict"] == "declined"]
    assert len(declines) == 1
    assert declines[0]["column"] is None and declines[0]["reason"] == summary["declined"]
    # Every artefact a mapped source has, so the review app lists it and the
    # corpus merge reads it — with the mapping document the one exception.
    for name in ("profile.json", "plan.json", "plan.validated.json",
                 "normalized.csv", "rejects.jsonl", "decisions.jsonl",
                 "entities.ftm.json", "statements.csv", "summary.json"):
        assert os.path.exists(os.path.join(d, name)), name
    assert _read(os.path.join(d, "entities.ftm.json")) == ""
    assert summary["mapping"] is None


def test_a_rerun_that_declines_publishes_no_mapping_at_all(tmp_path):
    """`ftm map <file>.yml` reproducing the entities without this pipeline is
    the claim the mapping document exists to support. A YAML left over from
    the run before, beside a summary saying no entity survived, would be two
    runs' answers in one directory.

    IT USED TO BE DELETED; NOW IT IS NEVER THERE. The mapping lives inside the
    generation it describes, so a declined rerun stages none and the CURRENT
    generation has none — while the previous generation keeps its own,
    correctly, because it really did produce those entities. A delete that has
    to run to keep the directory honest is a delete that can fail to run.
    """
    refs = inventory(_corpus(tmp_path))
    out = str(tmp_path / "work")
    mapped = run_source(refs[0], CFG, CAT, ScriptedClient(), "run1", out)
    first_mapping = mapped["mapping"]
    assert os.path.exists(first_mapping)

    declined = run_source(refs[0], CFG, CAT, NoEntitySurvives(), "run2", out)
    assert declined["declined"] == NO_ENTITY_REASON
    assert declined["mapping"] is None
    workspace = open_workspace(out)
    assert not os.path.exists(workspace.mapping_path(refs[0].safe_id))
    # ...and the superseded generation still holds the mapping it stood by.
    assert os.path.exists(first_mapping)
    assert first_mapping != workspace.mapping_path(refs[0].safe_id)


def test_a_workbook_is_parsed_once_per_run_not_once_per_sheet(tmp_path, monkeypatch):
    """`run_source` fetched its own grids, so a source — which is a SHEET —
    re-parsed the whole workbook and threw away every sheet but one.

    Measured on `data/files`, 22 workbooks expanding to 246 sources: parsing
    each workbook once costs 111s, once per sheet costs 1 755s, so 94% of all
    read time was spent building grids that were discarded. It stayed invisible
    on the public corpus, where 355 files expand to 527 sources — about 1.5
    sheets each — and shows up here only because these workbooks average 11.
    """
    import ftmap.pipeline as pipeline

    d = tmp_path / "corpus"
    d.mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "перший"
    wb.active.append(["ПІБ", "Дата народження"])
    wb.active.append(["Шевченко Ольга", "01.02.1990"])
    for name in ("другий", "третій"):
        ws = wb.create_sheet(name)
        ws.append(["ПІБ", "Дата народження"])
        ws.append(["Коваленко Іван Петрович", "17.09.1980"])
    wb.save(d / "b.xlsx")

    parses: list[str] = []
    real = pipeline.read_source

    def counting(path):
        parses.append(path)
        return real(path)

    monkeypatch.setattr(pipeline, "read_source", counting)
    report = run_corpus(str(d), CFG, CAT, ScriptedClient(), str(tmp_path / "out"))

    assert report["sources"] == 3, "three sheets is three sources"
    assert len(parses) == 1, f"one workbook, one parse; got {len(parses)}"


def test_two_workbooks_are_each_parsed_once(tmp_path, monkeypatch):
    """The cache holds ONE workbook, because a corpus does not fit in memory —
    `data/files` alone builds 56.8 million cells. `inventory` sorts by path and
    emits a file's sheets consecutively, so one is all that is ever needed."""
    import ftmap.pipeline as pipeline

    d = tmp_path / "corpus"
    d.mkdir()
    # DIFFERENT CONTENT PER FILE, or the second is correctly filtered as a
    # duplicate (`inventory` keys on sha256) and never parsed at all.
    for name, who in (("a.xlsx", "Шевченко Ольга"), ("b.xlsx", "Коваленко Іван")):
        wb = openpyxl.Workbook()
        wb.active.append(["ПІБ", "Дата народження"])
        wb.active.append([who, "01.02.1990"])
        wb.create_sheet("другий").append(["ПІБ", who])
        wb.save(d / name)

    parses: list[str] = []
    real = pipeline.read_source
    monkeypatch.setattr(pipeline, "read_source",
                        lambda p: (parses.append(p), real(p))[1])
    out = str(tmp_path / "out")
    run_corpus(str(d), CFG, CAT, ScriptedClient(), out)
    # Two parses for two workbooks, four sheets — and each one reads the
    # IMMUTABLE SNAPSHOT rather than the corpus file, so the digest recorded
    # beside the statements still describes the bytes they came from.
    assert len(parses) == 2, [os.path.basename(p) for p in parses]
    assert all(p.startswith(os.path.join(out, ".snapshots")) for p in parses), \
        parses


def _two_source_corpus(tmp_path):
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    # DIFFERENT CONTENT PER FILE. `inventory` deduplicates on sha256, so two
    # byte-identical workbooks are ONE source and a two-source fixture built
    # from copies quietly tests a one-source run.
    for name, value in (("a.xlsx", "Коваленко Іван"), ("b.xlsx", "Петренко Ольга")):
        wb = openpyxl.Workbook()
        wb.active.title = "арк"
        wb.active.append(["ПІБ"])
        wb.active.append([value])
        wb.save(corpus / name)
    return str(corpus)


def test_on_progress_reports_every_source_in_order(tmp_path):
    seen = []
    run_corpus(_two_source_corpus(tmp_path), CFG, CAT, ScriptedClient(),
               str(tmp_path / "work"), on_progress=seen.append)

    kinds = [event["kind"] for event in seen]
    assert kinds == ["source_started", "source_done"] * 2
    assert seen[0]["total"] == 2
    assert seen[0]["index"] == 0 and seen[2]["index"] == 1
    assert all(event["safe_id"] for event in seen)


def test_stop_ends_the_run_between_sources_and_the_manifest_says_so(tmp_path):
    out = str(tmp_path / "work")
    report = run_corpus(_two_source_corpus(tmp_path), CFG, CAT, ScriptedClient(),
                        out, stop=lambda: True)

    assert report["cancelled"] is True
    assert report["sources"] == 0
    with open(os.path.join(open_workspace(out).aggregate_dir(), "run.json"), encoding="utf-8") as fh:
        assert json.load(fh)["cancelled"] is True


def test_stop_after_the_first_source_keeps_that_source(tmp_path):
    """A cancelled run is a real run over fewer sources, not a discarded one."""
    seen = []

    def stop() -> bool:
        return len([e for e in seen if e["kind"] == "source_done"]) >= 1

    report = run_corpus(_two_source_corpus(tmp_path), CFG, CAT, ScriptedClient(),
                        str(tmp_path / "work"), on_progress=seen.append, stop=stop)
    assert report["cancelled"] is True
    assert report["sources"] == 1


def test_a_run_with_neither_hook_is_unchanged(tmp_path):
    corpus = tmp_path / "one"
    corpus.mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "арк"
    wb.active.append(["ПІБ"])
    wb.active.append(["Коваленко Іван"])
    wb.save(corpus / "a.xlsx")

    report = run_corpus(str(corpus), CFG, CAT, ScriptedClient(),
                        str(tmp_path / "work"))
    assert report["cancelled"] is False
    assert report["sources"] == 1


def test_the_manifest_says_when_the_run_ended(tmp_path):
    """`run.json` recorded `started_at` and nothing else, so every reader had
    to take the wall clock from artefact mtimes — including the published
    measurement record, which had to say so in a footnote."""
    import datetime
    import json as _json

    out = str(tmp_path / "work")
    fixtures = os.path.join(os.path.dirname(__file__), "fixtures")
    run_corpus(os.path.join(fixtures, "deputies_popolo.csv"), CFG, CAT,
               ScriptedClient(), out)
    with open(os.path.join(open_workspace(out).aggregate_dir(), "run.json"), encoding="utf-8") as fh:
        report = _json.load(fh)
    assert report["finished_at"], "a finished run says when it finished"
    start = datetime.datetime.fromisoformat(report["started_at"])
    end = datetime.datetime.fromisoformat(report["finished_at"])
    assert end >= start


def test_a_model_backed_run_records_its_revision_and_producing_model(tmp_path):
    """R11, through the pipeline rather than at the client.

    A manifest is the artefact a measurement is read from six months later, so
    the two claims that make a run attributable — the deployment it declared
    and the model that actually answered — have to survive the trip from the
    client into `run.json`.
    """
    class Provenanced(ScriptedClient):
        def model_provenance(self):
            return {"revision": "qwen3.6-35b-a3b-ud-q4-k-xl/llamacpp-b1234",
                    "producing_model": "Qwen3.6-35B-A3B"}

    report = run_corpus(_corpus(tmp_path), CFG, CAT, Provenanced(),
                        str(tmp_path / "work"))
    assert report["model_revision"] == "qwen3.6-35b-a3b-ud-q4-k-xl/llamacpp-b1234"
    assert report["model_id"] == "Qwen3.6-35B-A3B"


def _json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def test_every_sheet_and_every_rerun_reads_one_snapshot_of_the_workbook(tmp_path):
    """R12, end to end: one physical input, one immutable copy.

    `sha256` was recorded once at inventory and every later stage re-opened
    the ORIGINAL path, so a workbook edited between the two produced entities
    and statements attributed to a digest that no longer described them —
    with a manifest asserting that digest and nothing re-checking it.
    """
    from ftmap.fsutil import SnapshotStore

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "перший"
    wb.active.append(["ПІБ"])
    wb.active.append(["Коваленко Іван"])
    wb.create_sheet("другий").append(["ПІБ"])
    wb["другий"].append(["Шевченко Ольга"])
    wb.save(corpus / "a.xlsx")

    out = str(tmp_path / "work")
    first = run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    assert first["sources"] == 2

    stored = [p for p in (tmp_path / "work" / SnapshotStore.DIRNAME).rglob("*")
              if p.is_file()]
    assert len(stored) == 1, "two sheets of one workbook took two snapshots"

    # A rerun into the same workspace reuses it rather than rewriting a file
    # earlier generations are already reading.
    before = stored[0].stat().st_ino
    run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    assert stored[0].stat().st_ino == before


def test_editing_the_input_after_inventory_cannot_change_what_is_emitted(tmp_path):
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "a.csv").write_text("ПІБ\nКоваленко Іван\n", encoding="utf-8")
    out = str(tmp_path / "work")
    run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    safe_id = _report(out)["summaries"][0]
    with open(os.path.join(_sdir(out, safe_id), "statements.csv"), encoding="utf-8") as fh:
        before = fh.read()

    # The corpus file is replaced wholesale between runs. The snapshot for the
    # OLD digest is untouched, so the old generation still reads what it read.
    (corpus / "a.csv").write_text("ПІБ\nЗовсім інше\n", encoding="utf-8")
    from ftmap.fsutil import SnapshotStore
    store = SnapshotStore(out)
    digest = _json(os.path.join(_sdir(out, safe_id), "summary.json"))["source"]["sha256"]
    assert store.verify(digest, ".csv")
    with open(os.path.join(_sdir(out, safe_id), "statements.csv"), encoding="utf-8") as fh:
        assert fh.read() == before


def test_a_tampered_snapshot_stops_a_rerun_before_it_writes_anything(tmp_path):
    """Fail closed, and fail BEFORE the source's own output is replaced.

    A snapshot is immutable by contract and not by permission. If the check
    ran after `run_source` had begun overwriting the source directory, a
    mismatch would leave a half-replaced generation — which is the failure
    mode, not the fix for it.
    """
    from ftmap.fsutil import SnapshotMismatch, SnapshotStore
    from ftmap.inventory import inventory

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "a.csv").write_text("ПІБ\nКоваленко Іван\n", encoding="utf-8")
    out = str(tmp_path / "work")
    run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    safe_id = _report(out)["summaries"][0]
    entities = os.path.join(_sdir(out, safe_id), "entities.ftm.json")
    with open(entities, encoding="utf-8") as fh:
        before = fh.read()

    store = SnapshotStore(out)
    refs = inventory(str(corpus), snapshots=store)
    with open(refs[0].snapshot, "w", encoding="utf-8") as fh:
        fh.write("ПІБ\nпідмінено\n")

    with pytest.raises(SnapshotMismatch):
        run_source(refs[0], CFG, CAT, ScriptedClient(), "run2", out)
    with open(entities, encoding="utf-8") as fh:
        assert fh.read() == before


def test_a_full_corpus_run_applies_the_decisions_already_on_record(tmp_path):
    """Covers R6 and finding #8.

    Persisted overrides used to apply only to a per-source rerun from the
    review app. A full corpus run ignored them, so the export an analyst
    published after a week of review contained none of it — the work was on
    disk, in the log, and absent from the artefact.
    """
    from ftmap.review.store import DecisionStore

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "a.csv").write_text("ПІБ,Посада\nКоваленко Іван,депутат\n",
                                  encoding="utf-8")
    out = str(tmp_path / "work")
    run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)

    store = DecisionStore(out)
    safe_id = open_workspace(out).source_ids()[0]
    store.record(safe_id, "c1", "Person:position", "person", "analyst says so")

    run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    workspace = open_workspace(out)
    assert workspace.entry(safe_id).override_revision == 1
    with open(os.path.join(workspace.source_dir(safe_id), "statements.csv"),
              encoding="utf-8") as fh:
        assert "position" in fh.read()

    # THE MODEL-ONLY BASELINE IS A FRESH WORKSPACE, not a flag. An in-place
    # "ignore overrides" mode can be left on by accident and produces a
    # directory holding a mixture of analyst and model-only results with
    # nothing distinguishing them.
    baseline = str(tmp_path / "baseline")
    run_corpus(str(corpus), CFG, CAT, ScriptedClient(), baseline)
    fresh = open_workspace(baseline)
    with open(os.path.join(fresh.source_dir(fresh.source_ids()[0]),
                           "statements.csv"), encoding="utf-8") as fh:
        assert "position" not in fh.read()


def test_a_failed_rerun_leaves_the_corpus_staleness_alone(tmp_path):
    """A rerun that raises committed nothing, so nothing about the corpus
    changed — including whether it is stale."""
    from ftmap.review.store import DecisionStore

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "a.csv").write_text("ПІБ\nКоваленко Іван\n", encoding="utf-8")
    out = str(tmp_path / "work")
    run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    before = open_workspace(out)
    generation = before.entry(before.source_ids()[0]).generation
    assert before.stale_sources() == []

    class Boom(ScriptedClient):
        def complete(self, system, user, schema, max_tokens=None):
            raise RuntimeError("the model is gone")

    refs = inventory(str(corpus))
    with pytest.raises(RuntimeError):
        run_source(refs[0], CFG, CAT, Boom(), "run2", out)

    after = open_workspace(out)
    assert after.entry(after.source_ids()[0]).generation == generation
    assert after.stale_sources() == []
    assert after.publication_grade()


def test_publishing_a_failed_run_does_not_make_it_publication_grade(tmp_path):
    """The worst thing a publish could do, and it did it.

    `publish_corpus` passed `failures=[]`, `cancelled=False` and `client=None`
    into `_commit_aggregate`, which derives `incomplete` as
    `bool(cancelled or failures)`. So republishing a run that had lost a
    source to an exception rewrote it as complete, `publication_grade()` went
    True, and the model that produced the entities was replaced by the
    `not_applicable` claim reserved for runs that never had one.
    """
    from ftmap.pipeline import publish_corpus
    from ftmap.workspace import open_workspace

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "good.csv").write_text("код,ПІБ\nUA1,Коваленко Іван\n",
                                     encoding="utf-8")
    # A file the readers cannot open: one real failure, and the run says so.
    (corpus / "broken.xlsx").write_bytes(b"not a workbook at all")

    out = str(tmp_path / "work")
    report = run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    assert report["failed"] == 1, report["failures"]
    assert report["incomplete"] is True
    assert open_workspace(out).publication_grade() is False
    model_id, revision = report["model_id"], report["model_revision"]

    published = publish_corpus(out, CFG, CAT)

    assert published["incomplete"] is True, "a publish un-failed a source"
    assert published["failed"] == 1
    assert open_workspace(out).publication_grade() is False
    # And it still says which deployment produced what it summarises.
    assert published["model_id"] == model_id
    assert published["model_revision"] == revision


def test_republishing_a_run_that_had_a_duplicate_file(tmp_path):
    """A publish replays the previous run's OWN record of its duplicates.

    `_commit_aggregate` read `duplicate_sources` off `SourceRef` objects —
    `r.path`, `r.sheet`, `r.duplicate_of` — which is what `run_corpus` holds.
    `publish_corpus` does not: it merges what is committed and takes the list
    out of the prior `run.json`, where it is dicts. So any republish of a
    corpus holding two files with the same bytes raised `AttributeError:
    'dict' object has no attribute 'path'` — every source committed, every
    artefact on disk, and the corpus exports unbuildable. Nothing covered it,
    because every publish test used a corpus of distinct files.
    """
    from ftmap.pipeline import publish_corpus

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    text = "ПІБ,Посада\nКоваленко Іван,депутат\n"
    (corpus / "a.csv").write_text(text, encoding="utf-8")
    # The same bytes under another name — a backup copy, a symlinked export.
    # `inventory` keeps the first and marks this one `duplicate_of` it.
    (corpus / "b.csv").write_text(text, encoding="utf-8")

    out = str(tmp_path / "work")
    report = run_corpus(str(corpus), CFG, CAT, ScriptedClient(), out)
    assert report["duplicates"] == 1
    assert report["sources"] == 1, "the duplicate must not be run twice"
    recorded = report["duplicate_sources"]
    assert recorded[0]["path"].endswith("b.csv")
    assert recorded[0]["duplicate_of"].endswith("a.csv")

    published = publish_corpus(out, CFG, CAT)

    assert published["duplicates"] == 1
    assert published["duplicate_sources"] == recorded
    # And the manifest on disk says the same thing after the republish.
    assert _report(out)["duplicate_sources"] == recorded


def test_only_selects_sources_by_path_sheet_or_id_and_keeps_inventory_order():
    """`--only` is how one sheet under an etalon is rerun without the
    fourteen beside it. Matched on the names a person has in front of
    them, case-sensitively — a sheet name is data — and nothing given
    selects everything, so a caller that does not filter is unchanged."""
    from ftmap.inventory import SourceRef
    from ftmap.pipeline import select_sources

    def ref(path, sheet, sid):
        return SourceRef(path=path, sheet_index=0, sheet=sheet, sha256="a" * 64,
                         size=1, ext=".xlsx", source_id=sid)
    refs = [ref("/c/2 мср.xlsx", "2. Штат", "f83e/2. Штат"),
            ref("/c/2 мср.xlsx", "6. СОЧ", "f83e/6. СОЧ"),
            ref("/c/79 МСП.csv", "", "b29b/")]
    assert select_sources(refs, None) == refs
    assert select_sources(refs, []) == refs
    assert [r.sheet for r in select_sources(refs, ["Штат"])] == ["2. Штат"]
    assert [r.sheet for r in select_sources(refs, ["штат"])] == []
    assert [r.path for r in select_sources(refs, ["79 МСП", "b29b/"])] == ["/c/79 МСП.csv"]
    assert [r.sheet for r in select_sources(refs, ["СОЧ", "Штат"])] == ["2. Штат", "6. СОЧ"]
